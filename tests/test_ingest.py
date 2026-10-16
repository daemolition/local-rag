import re
from pathlib import Path
from types import SimpleNamespace

import pytest
from flask import Flask
from qdrant_client import QdrantClient
from werkzeug.security import generate_password_hash

from app.auth import access
from app.auth.user_admin_routes import users_bp
from app.core.models import User
from app.core.settings_service import SettingsService
from app.ingest import indexer
from app.rag import knowledge
from app.vector import ingestion as ingestion_module
from app.vector.collections import ensure_collection
from app.vector.ingestion import DocumentIngestion
from app.web.admin_routes import admin_bp
from app.web.api_routes import api_bp
from app.web.routes import bp as main_bp

VOCAB = ["kuendigung", "kaution", "miete", "dach", "frist", "strom"]


class FakeDense:
    dimension = len(VOCAB) + 1

    def __init__(self):
        self.calls = []

    def embed_documents(self, texts):
        self.calls.append(list(texts))
        return [[float(sum(w[:6] in t.lower() for _ in [0])) for w in VOCAB] + [0.01] for t in texts]

    def embed_query(self, text):
        return self.embed_documents([text])[0]


class FakeSparse:
    def embed_documents(self, texts):
        return [SimpleNamespace(indices=[i for i, w in enumerate(VOCAB) if w[:6] in t.lower()] or [99], values=[1.0] * (len([i for i, w in enumerate(VOCAB) if w[:6] in t.lower()]) or 1)) for t in texts]


DOC = """# Mietvertrag

Vorspann.

## 1 Miete

Die Kaltmiete betraegt 950 Euro monatlich und ist bis zum dritten Werktag zu zahlen. {pad}

## 2 Kaution

### 2.1 Hoehe

Die Kaution betraegt drei Monatsmieten. {pad}
""".format(pad="Weitere Hinweise zur Zahlung. " * 12)


# --- build_points ----------------------------------------------------------


def points(markdown=DOC, **kw):
    base = dict(
        filename="v.md", source="data/files/v.md", file_group_id="fg-1",
        dense=FakeDense(), sparse=FakeSparse(), chunk_size=400, chunk_overlap=50,
    )
    base.update(kw)
    return indexer.build_points(markdown, **base)


def test_points_carry_headers_order_and_both_vectors():
    pts = points(extra_metadata={"session_id": "s-1"})
    assert len(pts) >= 3
    p = pts[-1]
    assert p.payload["headers"] == ["Mietvertrag", "2 Kaution", "2.1 Hoehe"]
    assert p.payload["title"] == "2.1 Hoehe"
    assert (p.payload["filename"], p.payload["file_group_id"], p.payload["type"]) == ("v.md", "fg-1", "document_chunk")
    assert [x.payload["chunk_index"] for x in pts] == list(range(len(pts)))
    assert {x.payload["chunk_total"] for x in pts} == {len(pts)}
    assert all(x.payload["session_id"] == "s-1" for x in pts)
    assert set(p.vector) == {"", "langchain-sparse"}  # Dense UND BM25
    assert len({x.id for x in pts}) == len(pts)


def test_embedded_text_has_the_section_path_but_original_content_stays_raw():
    pts = points()
    leaf = next(p for p in pts if "### 2.1 Hoehe" in p.payload["original_content"])
    # Ueberschriften, die schon im Text stehen, werden nicht doppelt vorangestellt
    assert leaf.payload["page_content"].startswith("Mietvertrag\n\n## 2 Kaution\n### 2.1 Hoehe")
    assert "Mietvertrag >" not in leaf.payload["original_content"]
    # Fortsetzungs-Chunks (ohne eigene Ueberschrift) bekommen den vollen Pfad
    tail = [p for p in pts if p.payload["headers"][-1:] == ["1 Miete"] and not p.payload["original_content"].startswith("##")]
    assert tail and all(p.payload["page_content"].startswith("Mietvertrag > 1 Miete\n\n") for p in tail)
    assert not any(p.payload["original_content"].startswith(".") for p in pts)


def test_all_chunks_are_embedded_in_one_batch_and_the_path_is_what_gets_embedded():
    dense = FakeDense()
    pts = points(dense=dense)
    assert len(dense.calls) == 1 and len(dense.calls[0]) == len(pts)
    assert dense.calls[0] == [p.payload["page_content"] for p in pts]


def test_enrichment_result_replaces_the_embedded_text():
    seen = []

    def enrich(excerpt, text):
        seen.append((len(excerpt), text))
        return f"Kontext: X\n\nInhalt: {text}"

    pts = points(enrich=enrich)
    assert len(seen) == len(pts) and all(n <= 2500 for n, _ in seen)
    assert all(p.payload["page_content"].startswith("Kontext: X") for p in pts)
    assert not any(p.payload["original_content"].startswith("Kontext") for p in pts)


@pytest.mark.parametrize("md", ["", "  \n", "# Nur Ueberschrift\n"])
def test_markdown_without_content_gives_no_points(md):
    assert points(md) == []


def test_fixed_point_ids_make_results_reproducible():
    ids = [f"id-{i}" for i in range(50)]
    pts = points(point_ids=ids)
    assert [p.id for p in pts] == ids[: len(pts)]


def test_markdown_roundtrip(tmp_path):
    path = indexer.save_markdown(tmp_path / "md" / "1", "abc", "# Äöü Titel\n\nText")
    assert path == indexer.markdown_file(tmp_path / "md" / "1", "abc")
    assert path.read_text(encoding="utf-8") == "# Äöü Titel\n\nText"


# --- DocumentIngestion -----------------------------------------------------


class Shared:
    """In-Memory-Qdrant, das close() ueberlebt (DocumentIngestion schliesst seinen Client)."""

    def __init__(self, client):
        self._c = client

    def close(self):
        pass

    def __getattr__(self, name):
        return getattr(self._c, name)


@pytest.fixture
def qd(monkeypatch):
    client = QdrantClient(":memory:")
    shared = Shared(client)
    monkeypatch.setattr(ingestion_module, "get_qdrant_client", lambda **kw: shared)
    return client


def make_ingestion(qd, **kw):
    kw.setdefault("embeddings", FakeDense())
    kw.setdefault("sparse_embeddings", FakeSparse())
    ing = DocumentIngestion(**kw)
    ing._shared_sparse = kw["sparse_embeddings"]
    return ing


@pytest.fixture
def no_enrichment(db):
    s = db.get_session()
    settings = SettingsService(s)
    settings.set("INGEST_CONTEXT_ENRICHMENT", "false")
    settings.set("EMBEDDING_DIMENSION", "7")  # Dimension der Fake-Vektoren
    s.close()


def set_setting(db, key, value):
    session = db.get_session()
    SettingsService(session).set(key, value)
    session.close()


def put_file(name="v.md", text=DOC, folder="data/files"):
    path = Path(folder) / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def scroll_all(qd, collection="open_knowledge_ui"):
    points, _ = qd.scroll(collection, limit=200, with_payload=True, with_vectors=False)
    return sorted((p for p in points if p.payload.get("type") == "document_chunk"), key=lambda p: p.payload["chunk_index"])


def test_a_markdown_file_is_stored_split_by_headers_and_archived(db, qd, no_enrichment):
    set_setting(db, "INGEST_CHUNK_SIZE", "400")
    put_file()
    with make_ingestion(qd) as ing:
        ing.ingest_documents()

    chunks = scroll_all(qd)
    assert len(chunks) >= 3 and chunks[-1].payload["headers"][-1] == "2.1 Hoehe"

    docs = db.get_all_documents()
    assert len(docs) == 1
    doc = docs[0]
    assert doc.parser_version == indexer.PARSER_VERSION and doc.knowledge_id == knowledge.get_knowledge(None)["id"]
    md_file = Path(doc.markdown_path)
    assert md_file.name == f"{doc.document_id}.md" and md_file.parent == Path("data/markdown/1")
    assert md_file.read_text(encoding="utf-8") == DOC  # .md bleibt unveraendert
    assert not Path("data/files/v.md").exists() and Path("data/processed_files/v.md").exists()
    assert {c.payload["file_group_id"] for c in chunks} == {doc.document_id}


def test_chunk_size_comes_from_the_settings(db, qd, no_enrichment):
    set_setting(db, "INGEST_CHUNK_SIZE", "250")
    put_file()
    with make_ingestion(qd) as ing:
        ing.ingest_documents()
    assert max(len(c.payload["original_content"]) for c in scroll_all(qd)) <= 300
    assert len(scroll_all(qd)) > 4


def test_context_enrichment_follows_the_setting(db, qd, monkeypatch):
    calls = []
    monkeypatch.setattr(DocumentIngestion, "_enrich_with_context", lambda self, doc, chunk: calls.append(chunk) or f"Kontext: Z\n\nInhalt: {chunk}")
    put_file()
    with make_ingestion(qd) as ing:  # Default: an
        assert ing.enrich_context is True
        ing.ingest_documents()
    assert calls and all(c.payload["page_content"].startswith("Kontext: Z") for c in scroll_all(qd))


def test_attachment_mode_stores_no_markdown_and_no_document_row(db, qd, no_enrichment):
    ensure_collection(qd, "user_uploads", 7)
    path = put_file("anhang.md", folder="data/attachments/s-1/a-1")
    with make_ingestion(
        qd, collection_name="user_uploads", track_documents=False, enrich_context=False,
        extra_metadata={"session_id": "s-1"}, files_dir="data/attachments/s-1/a-1",
    ) as ing:
        assert ing._process_file(str(path), ing._shared_sparse) is True
    assert all(c.payload["session_id"] == "s-1" for c in scroll_all(qd, "user_uploads"))
    assert db.get_all_documents() == [] and not Path("data/markdown").exists()
    assert path.exists()  # Anhaenge werden nicht archiviert


def test_empty_files_are_skipped_without_a_document_row(db, qd, no_enrichment):
    put_file("leer.md", "  \n")
    put_file("nur-titel.md", "# Nur ein Titel\n")
    with make_ingestion(qd) as ing:
        ing.ingest_documents()
    assert db.get_all_documents() == [] and scroll_all(qd) == []


def test_other_knowledge_uses_its_own_markdown_folder(db, qd, no_enrichment):
    kb = knowledge.create_knowledge("Zweites")
    ensure_collection(qd, kb["collection_name"], 7)
    paths = knowledge.knowledge_paths(kb)
    put_file("z.md", folder=str(paths.files_dir))
    with make_ingestion(qd, knowledge=kb) as ing:
        ing.ingest_documents()
    doc = db.get_all_documents(knowledge_id=kb["id"])[0]
    assert Path(doc.markdown_path).parent == Path(f"data/markdown/{kb['id']}")
    knowledge.delete_knowledge(kb["id"])
    assert not paths.markdown_dir.exists()  # wird mit dem Knowledge geloescht


# --- Docling (echter Parser) -----------------------------------------------


@pytest.fixture
def docx_file(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.shared import Inches
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (300, 150), (200, 30, 30))
    ImageDraw.Draw(img).rectangle([20, 20, 280, 130], outline="white", width=6)
    img.save(tmp_path / "a.png")

    d = docx.Document()
    d.add_heading("Vertrag", level=1)
    d.add_paragraph("Allgemeiner Vorspann des Vertrags.")
    d.add_heading("1 Miete", level=2)
    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate([("Posten", "Betrag"), ("Kaltmiete", "950 Euro")]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.add_heading("2 Skizze", level=2)
    d.add_paragraph("Der Grundriss:")
    d.add_picture(str(tmp_path / "a.png"), width=Inches(2))
    d.add_paragraph("Ende.")
    path = tmp_path / "vertrag.docx"
    d.save(path)
    return path


class FakeVision:
    def __init__(self):
        self.calls = 0

    def generate_image_message(self, b64):
        return b64

    def generate(self, message):
        self.calls += 1
        return "Ein roter Rahmen\n# Kein Titel"


def test_docling_produces_markdown_with_headings_tables_and_image_descriptions(docx_file):
    from app.components import CustomDOCXLoader

    vision = FakeVision()
    docs = list(CustomDOCXLoader(str(docx_file), model=vision).lazy_load())
    assert len(docs) == 1
    md = docs[0].page_content
    assert docs[0].metadata["filename"] == "vertrag.docx"
    assert re.search(r"^#{1,3} 1 Miete$", md, re.M)
    assert "| Kaltmiete | 950 Euro |" in md.replace("  ", " ")  # Markdown-Tabelle
    assert vision.calls == 1
    # Beschreibung steht an der Stelle des Bildes, nach dem Satz davor
    assert md.index("Der Grundriss:") < md.index("[BILD-BESCHREIBUNG: Ein roter Rahmen") < md.index("Ende.")
    assert "\n# Kein Titel" not in md  # '#' in der Beschreibung ist entschaerft


def test_docx_ingestion_end_to_end_keeps_the_section_path(db, qd, no_enrichment, docx_file):
    set_setting(db, "INGEST_CHUNK_SIZE", "120")  # kleine Abschnitte nicht zusammenfassen
    put_file_path = Path("data/files") / docx_file.name
    put_file_path.parent.mkdir(parents=True, exist_ok=True)
    put_file_path.write_bytes(docx_file.read_bytes())

    with make_ingestion(qd) as ing:
        ing.model = FakeVision()
        ing.ingest_documents()

    chunks = scroll_all(qd)
    assert chunks and chunks[0].payload["filename"] == "vertrag.docx"
    skizze = next(c for c in chunks if "Grundriss" in c.payload["original_content"])
    assert skizze.payload["headers"][-1] == "2 Skizze" and "BILD-BESCHREIBUNG" in skizze.payload["original_content"]
    doc = db.get_all_documents()[0]
    assert "BILD-BESCHREIBUNG" in Path(doc.markdown_path).read_text(encoding="utf-8")


# --- API: Markdown ansehen -------------------------------------------------


def add_user(db, username, role):
    s = db.get_session()
    u = User(username=username, display_name=username, password_hash=generate_password_hash("pw"), role=role, is_active=True)
    s.add(u)
    s.commit()
    uid = u.id
    s.close()
    return {"id": uid, "role": role}


@pytest.fixture
def app(db):
    app = Flask(__name__)
    app.secret_key = "t"
    app.extensions["ingestion_status"] = {"is_ingesting": False, "knowledge_id": None}
    for bp in (main_bp, api_bp, admin_bp, users_bp):
        app.register_blueprint(bp)
    return app


def client_for(app, user):
    c = app.test_client()
    with c.session_transaction() as s:
        s["user_id"] = user["id"]
    return c


def stored_doc(db, kb, text="# Titel\n\nInhalt äöü.\n", file_group="fg-9"):
    path = indexer.save_markdown(Path("data/markdown") / str(kb["id"]), file_group, text)
    db.add_document(file_group, "bericht.pdf", knowledge_id=kb["id"], markdown_path=str(path), parser_version=1)
    return db.get_all_documents(knowledge_id=kb["id"])[0], path


def test_markdown_can_be_viewed_and_downloaded(app, db):
    kb = knowledge.get_knowledge(None)
    doc, _ = stored_doc(db, kb)
    editor = client_for(app, add_user(db, "ed", "editor"))  # Standard ist oeffentlich lesbar

    res = editor.get(f"/api/documents/{doc.id}/markdown")
    assert res.status_code == 200 and res.json == {"filename": "bericht.pdf", "markdown": "# Titel\n\nInhalt äöü.\n"}

    dl = editor.get(f"/api/documents/{doc.id}/markdown?download=1")
    assert dl.mimetype == "text/markdown" and "attachment" in dl.headers["Content-Disposition"]
    assert dl.headers["Content-Disposition"].endswith("bericht.md")
    assert dl.get_data(as_text=True) == "# Titel\n\nInhalt äöü.\n"


def test_markdown_respects_knowledge_access_and_role(app, db):
    private = knowledge.create_knowledge("Privat")
    doc, _ = stored_doc(db, private)
    stranger = client_for(app, add_user(db, "ed", "editor"))
    plain = client_for(app, add_user(db, "us", "user"))
    owner = add_user(db, "own", "editor")
    access.add_grant("knowledge", private["id"], "user", owner["id"], "read")

    assert stranger.get(f"/api/documents/{doc.id}/markdown").status_code == 404  # kein Zugriff
    assert plain.get(f"/api/documents/{doc.id}/markdown").status_code == 404  # kein Zugriff, auch als Benutzer
    assert client_for(app, owner).get(f"/api/documents/{doc.id}/markdown").status_code == 200
    assert client_for(app, {"id": 1}).get(f"/api/documents/{doc.id}/markdown").status_code == 404  # auch Admins nicht


def test_documents_without_markdown_say_so(app, db):
    kb = knowledge.get_knowledge(None)
    db.add_document("old", "alt.pdf", knowledge_id=kb["id"])
    doc = db.get_all_documents()[0]
    res = client_for(app, {"id": 1}).get(f"/api/documents/{doc.id}/markdown")
    assert res.status_code == 404 and "Neu-Indexieren" in res.json["error"]
    assert client_for(app, {"id": 1}).get("/api/documents/999/markdown").status_code == 404


def test_paths_outside_the_markdown_folder_are_never_served(app, db, tmp_path):
    kb = knowledge.get_knowledge(None)
    secret = tmp_path / "geheim.txt"
    secret.write_text("GEHEIM")
    db.add_document("evil", "x.pdf", knowledge_id=kb["id"], markdown_path=str(secret))
    doc = db.get_all_documents()[0]
    res = client_for(app, {"id": 1}).get(f"/api/documents/{doc.id}/markdown")
    assert res.status_code == 404 and "GEHEIM" not in res.get_data(as_text=True)


def test_deleting_a_document_removes_its_markdown_file(app, db, monkeypatch):
    kb = knowledge.get_knowledge(None)
    doc, path = stored_doc(db, kb)
    access.add_grant("knowledge", kb["id"], "user", 1, "write")  # Admins haben keinen Sonderzugriff
    assert path.exists()
    res = client_for(app, {"id": 1}).delete(f"/api/documents/{doc.id}")
    assert res.status_code == 200 and not path.exists()


def test_document_list_flags_which_documents_have_markdown(app, db):
    kb = knowledge.get_knowledge(None)
    stored_doc(db, kb)
    db.add_document("old", "alt.pdf", knowledge_id=kb["id"])
    flags = {d["filename"]: d["has_markdown"] for d in client_for(app, {"id": 1}).get("/api/documents").json}
    assert flags == {"bericht.pdf": True, "alt.pdf": False}


# --- Tabellen als CSV ------------------------------------------------------


def make_docx_with_tables(tmp_path, big_rows=15):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_heading("Preisliste", level=1)
    d.add_paragraph("Vorspann zur Preisliste.")
    d.add_heading("1 Kleine Tabelle", level=2)
    small = d.add_table(rows=3, cols=2)
    for r, row in enumerate([("Posten", "Betrag"), ("Miete", "950"), ("Kaution", "2850")]):
        for c, v in enumerate(row):
            small.cell(r, c).text = v
    d.add_heading("2 Grosse Tabelle", level=2)
    big = d.add_table(rows=big_rows + 1, cols=3)
    for c, v in enumerate(("Monat", "Umsatz", "Kosten")):
        big.cell(0, c).text = v
    for r in range(1, big_rows + 1):
        for c, v in enumerate((f"Monat{r:02d}", str(r * 100), str(r * 40))):
            big.cell(r, c).text = v
    d.add_heading("3 Layout", level=2)
    tiny = d.add_table(rows=1, cols=2)  # nur eine Zeile: Layout, keine Daten
    tiny.cell(0, 0).text = "Links"
    tiny.cell(0, 1).text = "Rechts"
    path = tmp_path / "preise.docx"
    d.save(path)
    return path


def ingest_docx(qd, path, **kw):
    target = Path("data/files") / path.name
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(path.read_bytes())
    with make_ingestion(qd, **kw) as ing:
        ing.model = FakeVision()
        ing.ingest_documents()


def test_tables_become_csv_files_and_references_in_the_markdown(db, qd, no_enrichment, tmp_path):
    ingest_docx(qd, make_docx_with_tables(tmp_path))

    csvs = sorted(p.name for p in Path("./data").glob("preise--*--tabelle-*.csv"))
    assert [c.rsplit("--", 1)[1] for c in csvs] == ["tabelle-1.csv", "tabelle-2.csv"]  # Layout-Tabelle fehlt

    import pandas as pd

    big = pd.read_csv(Path("./data") / csvs[1])
    assert list(big.columns) == ["Monat", "Umsatz", "Kosten"] and len(big) == 15
    assert big["Umsatz"].sum() == sum(r * 100 for r in range(1, 16))

    md = Path(db.get_all_documents()[0].markdown_path).read_text(encoding="utf-8")
    assert "| Miete" in md and f"[TABELLE 1: als Datei `{csvs[0]}`" in md  # klein: bleibt stehen
    assert f"[TABELLE 2: Datei `{csvs[1]}`, 15 Zeilen, Spalten: Monat, Umsatz, Kosten" in md  # gross: Verweis
    assert "run_pandas" in md and "Monat01" in md and "Monat15" not in md  # nur Vorschau
    assert "[TABELLE 3" not in md and "Links" in md  # Layout-Tabelle: kein CSV, Text bleibt


def test_extracted_tables_have_no_points_of_their_own_but_are_readable_by_the_pandas_tool(db, qd, no_enrichment, tmp_path, worker):
    from app.rag import retrieval
    from app.rag.retrieval import RetrievalScope
    from app.tools.custom_tools import CustomTools

    ingest_docx(qd, make_docx_with_tables(tmp_path))
    refs, _ = qd.scroll("open_knowledge_ui", limit=100, with_payload=True, with_vectors=False)
    assert not [p for p in refs if p.payload.get("type") == "data_reference"]  # gehoeren zum Dokument
    assert {p.payload["filename"] for p in refs} == {"tabellen.docx"} or all("--tabelle-" not in p.payload["filename"] for p in refs)
    paths0 = knowledge.knowledge_paths(knowledge.get_knowledge(None))
    ref_names = {f"{paths0.name_prefix}{p.name}" for p in Path(paths0.data_dir).glob("*--tabelle-*.csv")}
    assert len(ref_names) == 2

    paths = knowledge.knowledge_paths(knowledge.get_knowledge(None))
    retrieval.set_scope(RetrievalScope(("open_knowledge_ui",), None, data_dirs=((paths.name_prefix, str(paths.data_dir)),)))
    tools = CustomTools(llm=None, retriever=None, data_dir=str(paths.data_dir))
    name = sorted(ref_names)[1]
    assert name in tools.list_files()
    assert str(sum(r * 40 for r in range(1, 16))) in tools.run_pandas(name, "result = df['Kosten'].sum()")


def test_table_extraction_is_skipped_for_attachments(db, qd, no_enrichment, tmp_path):
    ensure_collection(qd, "user_uploads", 7)
    src = make_docx_with_tables(tmp_path)
    target = Path("data/attachments/s-1/a-1") / src.name
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(src.read_bytes())
    with make_ingestion(
        qd, collection_name="user_uploads", track_documents=False, enrich_context=False,
        extra_metadata={"session_id": "s-1"}, files_dir="data/attachments/s-1/a-1",
    ) as ing:
        ing.model = FakeVision()
        assert ing._process_file(str(target), ing._shared_sparse)
    assert list(Path("./data").glob("*--tabelle-*.csv")) == []
    text = " ".join(c.payload["original_content"] for c in scroll_all(qd, "user_uploads"))
    assert "Monat15" in text  # Anhang: Tabelle bleibt komplett im Text


def test_deleting_the_document_removes_its_csvs_and_reference_chunks(app, db, qd, no_enrichment, tmp_path):
    ingest_docx(qd, make_docx_with_tables(tmp_path))
    app.extensions["client"] = qd
    other = Path("./data/andere.csv")
    other.write_text("a,b\n1,2\n")
    doc = db.get_all_documents()[0]
    access.add_grant("knowledge", knowledge.get_knowledge(None)["id"], "user", 1, "write")  # Admins haben keinen Sonderzugriff
    assert len(list(Path("./data").glob("*--tabelle-*.csv"))) == 2

    assert client_for(app, {"id": 1}).delete(f"/api/documents/{doc.id}").status_code == 200

    assert list(Path("./data").glob("*--tabelle-*.csv")) == []
    assert other.exists()  # fremde CSVs bleiben
    refs, _ = qd.scroll("open_knowledge_ui", limit=100, with_payload=True, with_vectors=False)
    assert not [p for p in refs if "--tabelle-" in p.payload.get("filename", "")]
    assert scroll_all(qd) == []


def test_the_original_can_be_discarded_after_ingestion(db, qd, no_enrichment):
    set_setting(db, "INGEST_KEEP_ORIGINAL", "false")
    put_file()
    with make_ingestion(qd) as ing:
        ing.ingest_documents()
    assert not Path("data/files/v.md").exists() and not Path("data/processed_files/v.md").exists()
    doc = db.get_all_documents()[0]
    assert Path(doc.markdown_path).read_text(encoding="utf-8") == DOC  # das Markdown bleibt
    assert scroll_all(qd)


def test_table_helpers_name_files_uniquely_and_safely():
    from app.ingest import tables

    a = tables.table_filename("Mein Vertrag (final).pdf", "ab12cd34-0000", 3)
    assert a == "Mein_Vertrag_final_.pdf--ab12cd34--tabelle-3.csv"
    assert "/" not in tables.table_filename("../../etc/passwd", "ab12cd34-0000", 1)
    assert tables.table_filename("x", "ab12cd34-0000", 1) != tables.table_filename("x", "ff00ff00-0000", 1)
    assert tables.table_pattern("ab12cd34-0000") == "*--ab12cd34--tabelle-*.csv"


def test_reference_chunk_of_a_plain_data_file_describes_columns_and_content(tmp_path):
    import pandas as pd

    from app.vector.ingestion import describe_data_file

    path = tmp_path / "umsatz.csv"
    pd.DataFrame({"Monat": ["Jan", "Feb"], "Umsatz": [100.0, 250.5]}).to_csv(path, index=False)
    text = describe_data_file(str(path))
    assert "`Umsatz`: 100 bis 250.5" in text and "`Monat`: z. B. Jan, Feb" in text


def test_the_reference_in_the_document_carries_the_column_hints_of_its_table():
    import pandas as pd

    from app.ingest import tables

    df = pd.DataFrame({"Monat": [f"Monat{i:02d}" for i in range(1, 16)], "Umsatz": [i * 100.0 for i in range(1, 16)]})
    info = {"name": "knowledge_1/x--ab12cd34--tabelle-2.csv", "rows": 15, "columns": list(df.columns), "frame": df}
    text = tables.render_reference(info, 2, "")
    assert "`Umsatz`: 100 bis 1500" in text and "Monat01" in text


def test_extracted_tables_get_no_reference_point_and_old_ones_are_removed(db, qd, no_enrichment, tmp_path):
    from app.vector import ingestion

    folder = tmp_path / "ref"
    folder.mkdir()
    (folder / "vertrag--ab12cd34--tabelle-1.csv").write_text("a,b\n1,2\n3,4\n")
    (folder / "preise.csv").write_text("a,b\n1,2\n3,4\n")
    deleted = []
    orig = ingestion.delete_data_reference
    ingestion.delete_data_reference = lambda col, name: deleted.append(name)
    try:
        embeddings = type("E", (), {"embed_documents": lambda self, t: [[0.0] * 7 for _ in t]})()
        sparse = type("S", (), {"embed_documents": lambda self, t: [type("V", (), {"indices": [1], "values": [1.0]})() for _ in t]})()
        ensure_collection(qd, "reftest", 7)
        assert ingestion.index_data_references(embeddings, sparse, "reftest", str(folder), name_prefix="k_1/") == 1
    finally:
        ingestion.delete_data_reference = orig
    assert deleted == ["k_1/vertrag--ab12cd34--tabelle-1.csv"]
    pts, _ = qd.scroll("reftest", limit=10, with_payload=True, with_vectors=False)
    assert [p.payload["filename"] for p in pts] == ["k_1/preise.csv"]


def test_failing_enrichment_does_not_lose_the_document(db, qd, no_enrichment):
    put_file("a.md", DOC)
    with make_ingestion(qd, enrich_context=True) as ing:
        def boom(*a, **k):
            raise ConnectionError("Modell nicht erreichbar")

        ing._enrich_with_context = boom
        ing.ingest_documents()
        assert ing.processed == 1 and ing.errors == []
        assert any("Kontext-Anreicherung" in w and "a.md" in w for w in ing.warnings)
    points = scroll_all(qd)
    assert points and all(p.payload["page_content"] for p in points)


def test_enrichment_changes_the_embedded_text_but_not_the_original(db, qd, no_enrichment):
    put_file("a.md", DOC)
    with make_ingestion(qd, enrich_context=True) as ing:
        ing._enrich_with_context = lambda doc, chunk: f"Kontext: Vertrag\n\nInhalt: {chunk}"
        ing.ingest_documents()
    point = scroll_all(qd)[0].payload
    assert point["page_content"].startswith("Kontext: Vertrag")
    assert not point["original_content"].startswith("Kontext:")


def test_failed_file_is_reported_and_leaves_no_orphan_tables(db, qd, no_enrichment):
    put_file("kaputt.md", DOC)
    put_file("gut.md", DOC)
    with make_ingestion(qd) as ing:
        original = ingestion_module.build_points
        calls = []

        def flaky(markdown, **kw):
            calls.append(kw["filename"])
            if kw["filename"] == "kaputt.md":
                raise RuntimeError("Embedding-Modell nicht geladen")
            return original(markdown, **kw)

        ingestion_module.build_points = flaky
        try:
            ing.ingest_documents()
        finally:
            ingestion_module.build_points = original
    assert ing.processed == 1
    assert [e["file"] for e in ing.errors] == ["kaputt.md"]
    assert "Embedding-Modell" in ing.errors[0]["error"]
    assert Path("data/files/kaputt.md").exists()  # bleibt zum erneuten Einlesen im Eingang
    assert {p.payload["filename"] for p in scroll_all(qd)} == {"gut.md"}


def test_data_file_can_be_previewed_and_downloaded_but_not_outside_the_folder(db, qd):
    from tests.test_access import add_user, client_for

    ed = add_user(db, "ed", "editor")
    Path("data").mkdir(exist_ok=True)
    Path("data/umsatz.csv").write_text("Monat,Umsatz\nJan,100\nFeb,250\n", encoding="utf-8")
    Path("data/geheim.txt").write_text("nein", encoding="utf-8")
    flask_app = Flask(__name__)
    flask_app.secret_key = "t"
    flask_app.register_blueprint(api_bp)
    c = client_for(flask_app, ed)

    view = c.get("/api/data-files/umsatz.csv")
    assert view.status_code == 200
    body = view.json
    assert body["columns"] == ["Monat", "Umsatz"] and body["rows"][0] == ["Jan", "100"]
    assert body["total_rows"] == 2 and body["truncated"] is False

    raw = c.get("/api/data-files/umsatz.csv?download=1")
    assert raw.status_code == 200 and b"Jan,100" in raw.data
    assert "attachment" in raw.headers["Content-Disposition"]

    assert c.get("/api/data-files/geheim.txt").status_code == 404  # keine strukturierte Datei
    assert c.get("/api/data-files/..%2Fapp.db").status_code == 404
    # Rolle "Benutzer" darf lesen, was ihr gehoert oder freigegeben/oeffentlich ist (hier: oeffentlicher Standard)
    assert client_for(flask_app, add_user(db, "bob", "user")).get("/api/data-files/umsatz.csv").status_code == 200


def test_chunks_carry_document_metadata_and_the_uploader(db, qd, no_enrichment):
    from app.auth import access
    from tests.test_access import add_user, client_for

    ed = add_user(db, "ed", "editor")
    kb = knowledge.get_knowledge(None)
    access.add_grant("knowledge", kb["id"], "user", ed["id"], "write")
    flask_app = Flask(__name__)
    flask_app.secret_key = "t"
    flask_app.extensions["ingestion_status"] = {"is_ingesting": False, "knowledge_id": None}
    flask_app.register_blueprint(api_bp)
    import io

    r = client_for(flask_app, ed).post(
        "/api/documents/upload",
        data={"files": (io.BytesIO(DOC.encode()), "vertrag.md")},
        content_type="multipart/form-data",
    )
    assert r.status_code == 200
    with make_ingestion(qd) as ing:
        ing.ingest_documents()
    payload = scroll_all(qd)[0].payload
    assert payload["file_type"] == "md" and payload["knowledge_id"] == kb["id"]
    assert payload["uploaded_by"] == ed["id"] and payload["uploaded_at"].endswith("+00:00")
    assert payload["parser_version"] == indexer.PARSER_VERSION and "splitter_version" in payload
    assert "page_start" not in payload  # Markdown hat keine Seiten
    assert not list(Path("data/files").glob("*.uploader"))


def test_page_marks_replace_docling_page_breaks_only_when_consistent():
    from app.components.preprocess_base import PAGE_BREAK, apply_page_marks

    md = f"Seite eins\n\n{PAGE_BREAK}\n\nSeite zwei\n\n{PAGE_BREAK}\n\nSeite drei"
    out = apply_page_marks(md, [1, 2, 3])
    assert out.startswith("<!--p:1-->") and "<!--p:2-->" in out and "<!--p:3-->" in out
    assert PAGE_BREAK not in out
    # Anzahl passt nicht: lieber keine Seiten als falsche
    wrong = apply_page_marks(md, [1, 2])
    assert "<!--p:" not in wrong and PAGE_BREAK not in wrong
    assert "<!--p:" not in apply_page_marks("Nur eine Seite", [1])
